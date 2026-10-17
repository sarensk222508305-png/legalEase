"""TXT, DOCX and PDF formatting utilities."""

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


def _clean_title(text: str, fallback: str = "Legal Document") -> str:
    first = text.strip().splitlines()[0] if text.strip() else fallback
    return re.sub(r"^[#*\s]+", "", first).strip() or fallback


def format_docx(text: str, doc_type: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type)
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = "Times New Roman"

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            doc.add_paragraph("")
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(line.lstrip("# ").strip() if line.startswith("#") else line)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        if line.startswith("#"):
            run.bold = True
            run.font.size = Pt(14 if line.startswith("##") else 16)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("LegalEase — Educational Draft")
    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(9)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalPDF(FPDF):
    def __init__(self, title: str):
        super().__init__()
        self.doc_title = title
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 8, "LegalEase | AI Legal Document Generator", 0, 1, "C")
        self.line(10, 17, 200, 17)
        self.ln(3)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "", 8)
        self.cell(0, 8, f"LegalEase — Educational Draft | Page {self.page_no()}", 0, 0, "C")


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalPDF(doc_type)
    pdf.set_title(doc_type)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, doc_type, align="C")
    pdf.ln(3)

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            pdf.ln(4)
            continue
        if line.startswith("#"):
            heading = line.lstrip("# ").strip()
            pdf.set_font("Helvetica", "B", 12 if line.startswith("##") else 14)
            pdf.multi_cell(0, 7, heading)
        elif line.startswith("- "):
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, "• " + line[2:])
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, line)

    return bytes(pdf.output())


def format_txt(text: str) -> bytes:
    return text.encode("utf-8")
